"""Read only the already generated Matrix sheet; never rerun geometry checks."""
import json
import sys
from io import BytesIO
from pathlib import Path
import openpyxl

def read_matrix(path):
    book = openpyxl.load_workbook(BytesIO(Path(path).read_bytes()), read_only=True, data_only=True)
    try:
        if 'Matrix' not in book.sheetnames:
            raise ValueError('Excel cũ không có sheet Matrix; cần tải lại file annotation.')
        sheet = book['Matrix']
        if sheet.max_column > 2003 or sheet.max_row > 5003 or (sheet.max_column-3)*(sheet.max_row-3)>500000:
            raise ValueError('Ma trận vượt giới hạn 500.000 ô.')
        rows = sheet.iter_rows(min_row=3)
        header = next(rows)
        frames = [int(c.value) for c in header[3:]]
        if any(f<0 or f>100000000 for f in frames) or len(set(frames))!=len(frames):
            raise ValueError('Danh sách frame trong Excel không hợp lệ.')
        tracks = []
        for row in rows:
            if row[0].value is None:
                continue
            present = []
            for index, cell in enumerate(row[3:]):
                color = cell.fill.fgColor if cell.fill is not None else None
                if cell.value is not None or color is not None and color.type == 'rgb' and str(color.rgb)[-6:].upper()=='E2EFDA':
                    present.append(frames[index])
            tracks.append(dict(id=str(row[0].value),label=str(row[1].value or ''),present=present))
        return dict(version=1,frames=frames,tracks=tracks)
    finally:
        book.close()

if __name__=='__main__':
    try:
        print(json.dumps(read_matrix(sys.argv[1]),ensure_ascii=False,allow_nan=False))
    except Exception as error:
        print(str(error),file=sys.stderr)
        sys.exit(1)
