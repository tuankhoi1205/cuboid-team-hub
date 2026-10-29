"""Count stored cuboid keyframes, including explicit outside records. Never infer edits."""
import json
import re
import zipfile
from collections import defaultdict

LIMIT = 30 * 1024 * 1024

def read_json(source):
    if zipfile.is_zipfile(source):
        with zipfile.ZipFile(source) as archive:
            entries = [v for v in archive.infolist() if not v.is_dir() and (v.filename.replace('\\', '/').split('/')[-1] == 'annotations.json' or v.filename.replace('\\', '/').endswith('.json') and 'annotations' in v.filename.replace('\\', '/').split('/')[:-1])]
            if len(entries) != 1:
                raise ValueError('ZIP phải có đúng một JSON trong annotations/. Với API CVAT, tải riêng JSON của một Job.')
            if entries[0].file_size > LIMIT:
                raise ValueError('JSON vượt 30 MB.')
            with archive.open(entries[0]) as stream:
                content = stream.read(LIMIT + 1)
    else:
        with open(source, 'rb') as stream:
            content = stream.read(LIMIT + 1)
    if len(content) > LIMIT:
        raise ValueError('JSON vượt 30 MB.')
    return json.loads(content.decode('utf-8-sig'), parse_constant=lambda value: (_ for _ in ()).throw(ValueError('Số không hữu hạn: ' + value)))

def integer(value):
    if isinstance(value, bool) or not re.fullmatch(r'\d{1,9}', str(value)):
        return None
    return int(value)

def flag(value):
    # Do not treat the string "false" as true.
    if value is True or value == 1 or value == 'true':
        return True
    if value is False or value == 0 or value == 'false':
        return False
    return None

def count_dataset(data, offset=1, job_index=None):
    job_count = 1
    backup = isinstance(data, list)
    if backup:
        job_count = len(data)
        if not data or len(data) > 10000:
            raise ValueError('Danh sách Job trong backup không hợp lệ.')
        if job_index is None and job_count != 1:
            raise ValueError(f'Backup chứa {job_count} Job. Chọn thứ tự Job (1–{job_count}) theo bản backup, không nhập Job ID.')
        job_index = 1 if job_index is None else job_index
        if not 1 <= job_index <= job_count:
            raise ValueError('Thứ tự Job nằm ngoài bản backup.')
        data = data[job_index - 1]
    if not isinstance(data, dict):
        raise ValueError('JSON phải là một object.')
    ids, seen, issues, frame_ids = {}, {}, [], set()
    duplicates = 0
    count_issues = []
    def add(object_id, frame, keyframe, outside, label='', server_id=None):
        nonlocal duplicates
        if object_id is None or frame is None:
            issues.append('Có cuboid thiếu Object ID hoặc frame hợp lệ.')
            count_issues.append('Có cuboid thiếu Object ID hoặc frame hợp lệ.')
            return
        if object_id > 100000000 or frame > 100000000:
            raise ValueError('ID/frame vượt phạm vi xử lý.')
        signature = (keyframe, outside, str(label))
        pair = (object_id, frame)
        if pair in seen:
            duplicates += 1
            if seen[pair] != signature:
                issues.append('Bản ghi trùng ID/frame có cờ keyframe/outside/class mâu thuẫn.')
                count_issues.append('Bản ghi trùng ID/frame mâu thuẫn.')
            return
        seen[pair] = signature
        if len(seen) > 100000:
            raise ValueError('Tối đa 100.000 bản ghi cuboid cho một lần đếm.')
        frame_ids.add(frame)
        row = ids.setdefault(object_id, dict(id=object_id, label=str(label), serverId=server_id, cuboids=0, visible=0, outside=0, keyframes=0, interpolated=0, unknown=0, firstFrame=frame, lastFrame=frame))
        if row['label'] != str(label):
            count_issues.append('Cùng Object ID có nhiều class trong file.')
        row['cuboids'] += 1
        row['firstFrame'], row['lastFrame'] = min(row['firstFrame'], frame), max(row['lastFrame'], frame)
        if keyframe is True and outside is not None:
            row['keyframes'] += 1
            row['outside' if outside else 'visible'] += 1
        elif keyframe is False:
            row['interpolated'] += 1
        else:
            row['unknown'] += 1
            issues.append('Có bản ghi thiếu cờ keyframe hoặc outside hợp lệ; không tự đoán.')

    if isinstance(data.get('items'), list) and isinstance(data.get('categories'), dict):
        kind, complete_outside = 'datumaro', False
        mapping = 'track_id + ' + str(offset)
        labels = data['categories'].get('label', {}).get('labels', [])
        if len(data['items']) > 2000:
            raise ValueError('Tối đa 2.000 frame.')
        scope_frames = []
        for item in data['items']:
            if not isinstance(item, dict):
                raise ValueError('Frame không hợp lệ.')
            frame = integer((item.get('attr') or {}).get('frame'))
            if frame is None:
                match = re.search(r'(?:^|[_/])(\d+)$', str(item.get('id', '')))
                frame = integer(match.group(1)) if match else None
            if frame is None:
                count_issues.append('Frame thiếu số thứ tự hợp lệ.')
            else:
                scope_frames.append(frame)
            for ann in item.get('annotations', []):
                if ann.get('type') != 'cuboid_3d':
                    continue
                attrs = ann.get('attributes') or {}
                track = integer(attrs.get('track_id'))
                label_id = integer(ann.get('label_id'))
                label = labels[label_id].get('name', '') if label_id is not None and label_id < len(labels) else str(ann.get('label_id', ''))
                # Datumaro omits outside=false; absence is visible, but the source
                # as a whole can omit outside=true records and stays incomplete.
                add(track + offset if track is not None else None, frame, flag(attrs.get('keyframe')), flag(attrs.get('outside', False)), label)
        if len(set(scope_frames)) != len(scope_frames):
            count_issues.append('Trùng frame trong danh sách items.')
        issues.append('Datumaro có thể bỏ outside. Dùng file ban đầu làm mốc cố định và chỉ bổ sung ID mới; không lấy tổng file sau trừ tổng file trước.')
    elif isinstance(data.get('tracks'), list) and isinstance(data.get('shapes', []), list) and isinstance(data.get('tags', []), list):
        kind, complete_outside = 'cvat-job-json', True
        mapping = 'Thứ tự Object khi mở lại Job: tags, shapes, tracks; bắt đầu từ ' + str(offset) + '. Không dùng database ID làm Object ID.'
        tags, shapes, tracks = data.get('tags', []), data.get('shapes', []), data['tracks']
        if len(tags) + len(shapes) + len(tracks) > 10000:
            raise ValueError('Tối đa 10.000 Object trong JSON Job.')
        for index, shape in enumerate(shapes):
            if shape.get('type') == 'cuboid':
                add(len(tags) + index + offset, integer(shape.get('frame')), True, flag(shape.get('outside', False)), shape.get('label', shape.get('label_id', '')), shape.get('id'))
        for index, track in enumerate(tracks):
            object_id = len(tags) + len(shapes) + index + offset
            if not isinstance(track, dict) or not isinstance(track.get('shapes'), list):
                raise ValueError('Track phải chứa danh sách shapes.')
            for shape in track['shapes']:
                if shape.get('type') != 'cuboid':
                    continue
                # Raw CVAT stores keyframes only; interpolated shapes are not in
                # this API representation. An explicit false flag is respected.
                add(object_id, integer(shape.get('frame')), flag(shape.get('keyframe', True)), flag(shape.get('outside')), track.get('label', track.get('label_id', '')), track.get('id'))
    else:
        raise ValueError('Cần Datumaro 3D hoặc annotations.json trong bản Backup CVAT có tracks và shapes. Không dùng metadata.')
    if not ids and not (kind == 'datumaro' and scope_frames):
        raise ValueError('Không tìm thấy cuboid có ID/frame hợp lệ để đếm.')
    if len(ids) > 5000:
        raise ValueError('Tối đa 5.000 Object ID cho một lần đếm.')
    rows = sorted(ids.values(), key=lambda row: row['id'])
    totals = {key: sum(row[key] for row in rows) for key in ['keyframes', 'visible', 'outside', 'interpolated', 'unknown']}
    notes = list(dict.fromkeys(issues))
    return dict(version=1, format='cvat-backup-json' if backup else kind, offset=offset, jobIndex=job_index, jobCount=job_count, completeOutside=complete_outside, usable=complete_outside and not notes, mapping=mapping, issues=notes, duplicates=duplicates, frames=len(frame_ids), totals=totals, ids=rows,
                datumaroUsable=kind == 'datumaro' and not count_issues, countIssues=list(dict.fromkeys(count_issues)),
                scopeFrames=sorted(set(scope_frames)) if kind == 'datumaro' else [], cuboidTotal=sum(r['cuboids'] for r in rows))

def add_count_sheet(workbook, production):
    from openpyxl.styles import Font, PatternFill, Alignment
    sheet = workbook.create_sheet('Cuboid theo ID')
    sheet.append(['Object ID', 'Class / Label ID', 'Keyframe hiện', 'Keyframe outside', 'Tổng cuboid keyframe', 'Nội suy không tính', 'Thiếu cờ', 'Frame đầu', 'Frame cuối', 'Database ID (không phải Object ID)'])
    for row in production['ids']:
        sheet.append([row['id'], row['label'], row['visible'], row['outside'], row['keyframes'], row['interpolated'], row['unknown'], row['firstFrame'], row['lastFrame'], row['serverId']])
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill, cell.font, cell.alignment = PatternFill('solid', fgColor='24417B'), Font(color='FFFFFF', bold=True), Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 42
    for column, width in zip('ABCDEFGHIJ', [15, 24, 18, 20, 25, 22, 16, 16, 16, 35]):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='EEF3FF')
            if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@', '\t', '\r')):
                cell.value = "'" + cell.value
    notes = workbook.create_sheet('Luu y KPI')
    for row in [['Nội dung', 'Chi tiết'], ['Định dạng', production['format']], ['Đếm đủ outside', 'Có bản ghi gốc' if production['completeOutside'] else 'Chưa đủ: Datumaro có thể bỏ outside'], ['Object ID', production['mapping']], ['Cách tính', 'Mỗi cặp Object ID/frame keyframe tính 1; bao gồm outside. Không tính nội suy, không nhân 0,85.'], ['Phạm vi ngày', 'Chỉ cộng ID lớn hơn mốc đã báo cáo các ngày trước và nhỏ hơn hoặc bằng mốc hôm nay.'], ['Giới hạn', 'Keyframe là khối lượng của ID được khai hoàn thành, không chứng minh ai đã chỉnh box prelabel.'], *[['Cần kiểm tra', note] for note in production['issues']]]:
        notes.append(row)
    notes.column_dimensions['A'].width, notes.column_dimensions['B'].width = 25, 115
    for row in notes:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        notes.row_dimensions[row[0].row].height = 42
    notes.freeze_panes = 'B2'
