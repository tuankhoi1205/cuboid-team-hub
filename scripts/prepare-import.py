"""Read-only conversion of the two user-provided Google Sheets exports.
No workbook is modified. Raw values, exact cells and source links are retained.
"""
import openpyxl, json, pathlib
root = pathlib.Path(__file__).resolve().parent.parent
src = root / '.sites-runtime/imports'
w = openpyxl.load_workbook(src/'workbook.xlsx', data_only=True)
k = openpyxl.load_workbook(src/'kpi.xlsx', data_only=True)
links = ['https://docs.google.com/spreadsheets/d/1tB3LUSJmHs6ZwZ1V5lKsyUEtoZQUnY3i/edit?gid=1078651171', 'https://docs.google.com/spreadsheets/d/1QpNIXP0umu9pTgd6-JiX_uJRrx_OAzMq/edit?gid=905526587']
records = []
project_id = 'import-project-cuboid'
def add(id,kind,data,**meta):
 records.append(dict(id=id,kind=kind,data=data,**meta))
add(project_id,'project',dict(name='3D Cuboid · KPI hiện tại',description='Nhập từ KPI_Ngay. Tên dự án gốc chưa được ghi trong nguồn; Lead có thể đổi tên.',instructions='Theo hướng dẫn annotation của team. Chỉ tiêu gốc: 1.200 cuboid/ngày. QA/QC phân bổ 50% annotation và 50% QC. Sản lượng lịch sử là số cuboid quy đổi ID × frame theo công thức trong sheet, chưa phải kết quả nghiệm thu QC.',unit='cuboid',target=1200,targetPeriod='day',source=links[0],sourceCell="'Bao cao tuan'!B4",imported=True),status='active')
s=k['Tuan1']; totals={}; bytask={}; tasks={}; issues=[]
for ri in range(3,12):
 member='member-'+str(ri-2); last_task=None
 for col in (2,7):
  day=s.cell(1,col).value.date().isoformat(); ids=s.cell(ri,col).value; frames=s.cell(ri,col+1).value; quantity=s.cell(ri,col+2).value; task=s.cell(ri,col+3).value; issue=s.cell(ri,col+4).value or ''
  if task is not None: last_task=str(task).removesuffix('.0')
  if ids is None or quantity is None: continue
  task_id='import-task-'+last_task; tasks.setdefault(last_task,dict(owner=member,day=day,frames=frames)); bytask[last_task]=bytask.get(last_task,0)+quantity
  line=dict(taskId=task_id,activity='annotation',minutes=None,quantity=quantity,rework=None,quantityClassification='net_daily_confirmed_by_lead',work='Sản lượng theo KPI_Ngay (ID × frame)',issue=issue,evidence='',sourceIds=ids,sourceFrames=frames,sourceTaskId=task,taskIdInherited=task is None,sourceRange=f"Tuan1!{openpyxl.utils.get_column_letter(col)}{ri}:{openpyxl.utils.get_column_letter(col+4)}{ri}")
  add(f'import-daily-{member}-{day}','daily',dict(date=day,lines=[line],completed='',blockers=issue if issue.lower() not in ('hiện chưa có','hôm nay chưa có') else '',support='',nextPlan='',imported=True,source=links[1],sourceState='Chưa duyệt',missingFields=['Thời gian thực tế'],quantityMethod='Giá trị sản lượng ngày trong sheet; ngày 21/10 dùng (ID cộng dồn - ID ngày trước) × số frame, không cộng ID cộng dồn trực tiếp.'),owner_id=member,date=day,natural_key=member+':'+day,status='imported')
  if issue and issue.lower() not in ('hiện chưa có','hôm nay chưa có'): issues.append(dict(memberId=member,date=day,text=issue,taskId=task_id))
  totals[day]=totals.get(day,0)+quantity
for code,t in tasks.items():
 add('import-task-'+code,'task',dict(code=code,title='Annotation 3D Cuboid · Task '+code,projectId=project_id,assigneeId=t['owner'],reviewerId='member-0',link='https://cvat-test.vantixapps.net/tasks/'+code,assignedDate=t['day'],assignedDateEstimated=True,deadline=None,priority='normal',quantity=None,completed=bytask[code],reportedTotal=bytask[code],latestProductionDate='2026-10-21',note='Dữ liệu nhập từ sheet. Chưa có khối lượng giao và deadline. Lead cần bổ sung trước khi tiếp tục workflow. Người kiểm tra tạm giao Lead vì chưa chọn QA/QC.',imported=True,source=links[1]),project_id=project_id,natural_key=code+':',status='doing')
s=w['Diem danh va tac phong']; statuses={'Đúng giờ':'present','Đi trễ':'late','Nghỉ có phép':'leave'}
for ri in range(3,12):
 member='member-'+str(ri-2)
 for col,day in [(2,'2026-10-20'),(6,'2026-10-21'),(10,'2026-10-22')]:
  status=s.cell(ri,col).value
  if status is None: continue
  add(f'import-attendance-{member}-{day}','attendance',dict(checkIn=None,checkOut=None,note=s.cell(ri,col+3).value or '',lateMinutes=s.cell(ri,col+1).value or 0,conduct=s.cell(ri,col+2).value,imported=True,source=links[0],sourceRange=f"'Diem danh va tac phong'!{openpyxl.utils.get_column_letter(col)}{ri}",originalDate=str(s.cell(1,col).value.date()),dateCorrection='Ngày 21/11 và 22/11 được người dùng xác nhận sửa thành 21/10 và 22/10.' if col>2 else None),owner_id=member,date=day,natural_key=member+':'+day,status=statuses[status])
add('import-undated-weekly','legacy',dict(title='2 ghi nhận rework chưa có tuần',source=links[0],rows=[dict(memberId='member-4',cell="'Bao cao tuan'!L11",reworkTasks=1),dict(memberId='member-5',cell="'Bao cao tuan'!L12",reworkTasks=1)],note='Nguồn chưa điền ngày bắt đầu tuần; giữ riêng, chưa cộng vào weekly report.'),status='needs_date')
add('import-manifest','import',dict(title='Nhập dữ liệu từ Google Sheets',sources=links,totals=totals,reportCount=18,taskCount=len(tasks),attendanceCount=27,notes=['Cầm Vũ Ngọc Thạch → Cẩm Vũ Ngọc Thạch; Phạm Nguyễn Tuấn → Phạm Nguyễn Tuân theo danh sách nhân sự người dùng.', 'Ngày điểm danh 21/11, 22/11 → 21/10, 22/10 theo xác nhận người dùng.', 'Tổng sản lượng từng ngày đã được Lead xác nhận trừ phần lỗi. QC thường vào sáng hôm sau; không tự tạo kết quả QC.', 'Báo cáo lịch sử chưa có thời gian thực tế, chưa được duyệt.', 'Không tạo dữ liệu cho các ô trống hoặc các tuần tương lai.', '2 ghi nhận rework thiếu ngày tuần được giữ riêng.', 'Chưa nhập điểm Quality vì không có mẫu số kiểm tra.'],issues=issues),status='completed')
out=root/'data/initial-import.json'; out.parent.mkdir(exist_ok=True); out.write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(records=len(records),tasks=len(tasks),daily=18,attendance=27,totals=totals),ensure_ascii=False))



